from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml import parse_xml
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
import json

ROOT=Path(__file__).parent;A=ROOT/'assets'
P='http://schemas.openxmlformats.org/presentationml/2006/main';AA='http://schemas.openxmlformats.org/drawingml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
Q=json.loads((ROOT/'questions.json').read_text(encoding='utf-8'))
assert len(Q)==8, 'This layout expects eight questions'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='لعبة كراش التفاعلية';prs.core_properties.subject='إعادة بناء مستوحاة من مرجعَي الفيديو';prs.core_properties.author='Reference Game Builder';prs.core_properties.keywords='Crash, Arabic, classroom, quiz, editable, interactive'
prs.core_properties.last_modified_by='Reference Game Builder'
SL={};pending=[];sound_parts={}
def rgb(s):return RGBColor.from_string(s)
def native_rect(s,x,y,w,h,fill,line=None,round=False,opacity=100):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background() if line is None else None
    if line:sh.line.color.rgb=rgb(line)
    if opacity!=100:
        col=sh._element.find('.//{'+AA+'}solidFill/{'+AA+'}srgbClr')
        el=OxmlElement('a:alpha');el.set('val',str(opacity*1000));col.append(el)
    if round:sh.adjustments[0]=.12
    return sh

def text(s,txt,x,y,w,h,size=28,color='3C1F0A',bold=False,align=PP_ALIGN.CENTER,name=None,font='Arial',outline=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name=name or 'TEXT'
    tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_top=tf.margin_bottom=0;tf.margin_left=tf.margin_right=Inches(.06);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_after=Pt(2);pPr=p._p.get_or_add_pPr();pPr.set('rtl','0' if txt.isascii() else '1')
        run=p.add_run();run.text=line;run.font.name=font;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)
        rp=run._r.get_or_add_rPr();rp.set('lang','en-US' if txt.isascii() else 'ar-SA')
        for tag in ['a:cs','a:ea']:
            el=OxmlElement(tag);el.set('typeface',font);rp.append(el)
        if outline:
            ln=OxmlElement('a:ln');ln.set('w','22000');fill=OxmlElement('a:solidFill');c=OxmlElement('a:srgbClr');c.set('val',outline);fill.append(c);ln.append(fill);rp.insert(0,ln)
    return sh

def pic(s,n,x,y,w,h=None,name=None):
    sh=s.shapes.add_picture(str(A/n),Inches(x),Inches(y),width=Inches(w),height=Inches(h) if h else None);sh.name=name or n;return sh

def link(shape,target):pending.append((shape,target))
def button(s,label,x,y,w,h,target,name=None):
    b=native_rect(s,x,y,w,h,'532509','EAB363',True);b.name=name or 'BUTTON';link(b,target)
    t=text(s,label,x+.04,y+.03,w-.08,h-.06,16 if h<.65 else 21,'FFF0BB',True);link(t,target);return b

def transition(s,effect='fade',auto=None,sound=None):
    t=OxmlElement('p:transition');t.set('spd','fast');t.set('advClick','0')
    if auto is not None:t.set('advTm',str(auto))
    e=OxmlElement('p:'+effect);t.append(e)
    if sound:
        if sound not in sound_parts:sound_parts[sound]=Part(PackURI('/ppt/media/'+sound+'.wav'),'audio/wav',prs.part.package,(A/(sound+'.wav')).read_bytes())
        rid=s.part.relate_to(sound_parts[sound],RT.AUDIO)
        action=OxmlElement('p:sndAc');start=OxmlElement('p:stSnd');start.set('loop','0');snd=OxmlElement('p:snd');snd.set('{'+R+'}embed',rid);snd.set('name',sound+'.wav');start.append(snd);action.append(start);t.append(action)
    s._element.insert(2,t)

def fade(s,shapes,stagger=80,duration=350):
    # Native PowerPoint after-previous animation timing, not an HTML animation.
    parts=[]
    for i,sh in enumerate(shapes):
        k=4+i*4;delay=i*stagger
        parts.append(f'''<p:par><p:cTn id="{k}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{k+1}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{k+2}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{k+3}" dur="{duration}"/><p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>''')
    xml=f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{''.join(parts)}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    s._element.append(parse_xml(xml))

def new(name,bg=None,effect='fade',auto=None,sound=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.name=name;SL[name]=s
    if bg:pic(s,bg,0,0,13.333333,7.5,'BACKGROUND')
    transition(s,effect,auto,sound);return s

def notes(s,msg):s.notes_slide.notes_text_frame.text=msg

def movie_slide(name,clip,target):
    s=new(name,auto=4100)
    m=s.shapes.add_movie(str(A/(clip+'.mp4')),0,0,prs.slide_width,prs.slide_height,poster_frame_image=str(A/(clip+'-poster.png')),mime_type='video/mp4')
    m.name='EMBEDDED_'+clip.upper()
    for cond in s._element.findall('.//{'+P+'}video/{'+P+'}cMediaNode/{'+P+'}cTn/{'+P+'}stCondLst/{'+P+'}cond'):cond.set('delay','0')
    button(s,'تخطّي المقدمة' if name=='INTRO' else 'إلى الصناديق',10.65,6.69,2.05,.5,target)
    if name=='PORTAL':
        text(s,'إلى عالم كراش',3.7,.28,6,.55,27,'FFF2C1',True)
    notes(s,'فيديو أصلي مدمج داخل الملف، مولّد بتصيير ثلاثي الأبعاد. يبدأ تلقائياً ثم ينتقل للصفحة التالية. زر التخطي يعمل حتى إذا منع عارض آخر تشغيل الفيديو.')
    return s

movie_slide('INTRO','intro','TITLE')
s=new('TITLE','beach-clean.jpg')
native_rect(s,0,0,13.333,7.5,'1C3219',opacity=18)
pic(s,'crash.png',.35,1.15,5.6,5.1)
# Native editable title styled after the reference, not a screenshot or logo image.
text(s,'CRASH',5.6,.95,7.05,1.5,91,'5B2F11',True,font='Impact')
title=text(s,'CRASH',5.54,.84,7.05,1.5,91,'FFE351',True,font='Impact',outline='6B350D')
text(s,'BANDICOOT',6.25,2.13,5.8,.64,36,'FFEAC3',True,font='Impact',outline='4F2A0E')
pic(s,'sign.png',5.7,2.55,6.9,2)
text(s,'لعبة كراش التفاعلية',6.4,3.17,5.5,.65,32,'482107',True)
button(s,'هيا بنا للإجابة على الأسئلة',6.4,4.78,5.5,.82,'PORTAL')
text(s,'٨ صناديق  •  أسئلة قابلة للتعديل',6.2,5.95,5.9,.45,20,'FFF8D8',True)
button(s,'طريقة اللعب والتعديل',9.56,6.71,2.85,.49,'TEACHER')
fade(s,[title],duration=700)
notes(s,'شغّل عرض الشرائح من البداية. بعد المقدمة اضغط هيا بنا. الملف يعتمد على روابط داخلية، لا يحتاج إنترنت أو وحدات ماكرو.')
movie_slide('PORTAL','portal','MENU')
s=new('MENU','beach.jpg')
pic(s,'sign.png',3.3,-.02,7.6,2.02)
text(s,'اختر الصندوق يا بطل',4.0,.64,6.2,.79,34,'FFF4D7',True,outline='794219')
crates=[]
for i in range(8):
    x=1.3+(i%4)*2.48;y=2.3+(i//4)*2.14
    c=pic(s,'crate.png',x,y,1.6,1.53,f'CRATE_{i+1}');link(c,f'Q{i+1:02d}');crates.append(c)
    label=text(s,str(i+1),x+.49,y+1.48,.62,.38,22,'FFFFFF',True,outline='6B350D');link(label,f'Q{i+1:02d}')
home=pic(s,'fruit.png',.35,.35,.72,.79,'HOME');link(home,'TITLE')
button(s,'إنهاء الجولة',.44,6.91,1.95,.42,'FINISH')
fade(s,crates,95,320)
notes(s,'اختر أي صندوق. كل صندوق مرتبط بسؤال. ثمرة الوُمبا تعيدك إلى البداية. زر إنهاء الجولة ينتقل إلى الختام. لا تُحسب نقاط تلقائية ولا تُحفظ حالة الصناديق، كما في قالب باوربوينت تقليدي بلا ماكرو.')

for i,q in enumerate(Q):
    base=f'Q{i+1:02d}'
    for selected in [None,*range(len(q['options']))]:
        correct=selected==q['correct'] if selected is not None else None
        s=new(base if selected is None else base+f'_A{selected+1}','jungle.jpg',effect='cut' if selected is not None else 'fade',sound=None if selected is None else 'correct' if correct else 'wrong')
        if selected is not None:
            native_rect(s,0,0,13.333,7.5,'295E19' if correct else '671820',opacity=13 if correct else 22)
        board=pic(s,'sign.png',3.95,-.02,8.9,2.7,'QUESTION_BOARD')
        question=text(s,q['text'],4.67,.90,7.45,1.37,30 if len(q['text'])<65 else 27,'3B1D08',True,name='QUESTION')
        text(s,q['category']+'  |  '+str(i+1)+' من ٨',.33,.26,3.3,.46,17,'FFF3C5',True,outline='57431B')
        answer_items=[]
        for j,opt in enumerate(q['options']):
            y=3.10+j*1.05 if len(q['options'])==3 else 3.38+j*1.30
            p=pic(s,'plank.png',4.73,y,7.8,.86,f'ANSWER_{j+1}');link(p,base+f'_A{j+1}');answer_items.append(p)
            t=text(s,opt,5.37,y+.08,6.43,.61,30,'42210D',True,name=f'OPTION_TEXT_{j+1}');link(t,base+f'_A{j+1}');answer_items.append(t)
            if j==selected:
                feedback=pic(s,'correct.png' if correct else 'wrong.png',3.92,y+.06,.74,.74,'CORRECT' if correct else 'WRONG')
        home=pic(s,'fruit.png',7.61,6.38,.78,.85,'HOME');link(home,'MENU')
        text(s,'الصناديق',8.42,6.68,1.6,.35,17,'402815',True)
        if selected is not None:
            status='أحسنت يا بطل!' if correct else 'حاول مرة أخرى'
            msg=text(s,status,.22,6.43,3.44,.65,27,'FFF0AC',True,outline='3E2D13')
            # Brief explanation stays on a separate teacher-friendly reveal slide state.
            if correct:
                text(s,q['explanation'],4.47,6.01,8.25,.4,16,'3F270F',True,name='EXPLANATION')
            fade(s,[feedback,msg],60,260)
        else:fade(s,[board,question,*answer_items],65,320)
        notes(s,f"السؤال {i+1}: {q['text']}\nالصحيح: {q['options'][q['correct']]}\n{q['explanation']}\nالتعديل: لأن التغذية الراجعة تعتمد على شرائح حالات مرتبطة لا على ماكرو، عدّل النص في شريحة {base} وشرائح {base}_A التابعة لها. روابط الخيارات ترتبط بمواضع الإجابات، فلا تنقل الصحيح دون تعديل الروابط/حالات التصحيح. يمكنك استبدال نصوص الخيار الصحيح والخاطئ مع الحفاظ على مواقعها. اضغط ثمرة الوُمبا للعودة للصناديق.")

s=new('FINISH','beach-clean.jpg');native_rect(s,0,0,13.333,7.5,'243720',opacity=17);pic(s,'crash.png',.44,1.45,5.3,4.86);pic(s,'sign.png',5.03,.57,7.85,3.11)
title=text(s,'شكراً على جهودكم\nيا أبطال!',5.88,1.63,6.15,1.65,39,'47220C',True)
button(s,'العبوا مرة أخرى',6.25,4.45,5,.79,'MENU');button(s,'العودة للبداية',7.2,5.56,3.1,.57,'TITLE');fade(s,[title],duration=600)
notes(s,'ختام تفاعلي مستوحى من المرجع الثاني. لا يعلن فائزاً أو نقاطاً غير محسوبة.')
s=new('TEACHER','beach.jpg');native_rect(s,.55,.35,12.25,6.8,'302211',line='E9B966',round=True,opacity=94)
text(s,'دليل المعلّم',1.1,.68,11.1,.72,35,'FFE096',True)
lines=[('التشغيل','ابدأ عرض الشرائح من أول شريحة. المقدمة والبوابة تعملان تلقائياً، ويمكن تخطيهما.'),('اللعب','ثمانية صناديق. اختر صندوقاً، ثم إجابة. علامة صح أو خطأ وصوت يوضّحان النتيجة.'),('التنقّل','ثمرة الوُمبا تعيدك إلى الصناديق. لا تضغط خارج الأزرار؛ التنقّل العشوائي معطّل.'),('تعديل الأسئلة','النصوص عربية وقابلة للتحرير. عدّل السؤال وخياراته في شريحة السؤال وشرائح النتائج التابعة لها.'),('نظام الإجابات','احتفظ بالخيار الصحيح في موضعه عند استبدال النص. توجد إرشادات ومفتاح الإجابة في الملاحظات.'),('التوافق','النسخة مخصّصة لباوربوينت سطح المكتب. الفيديوهات والأصوات مدمجة. لا تحتاج إنترنت أو ماكرو.')]
for k,(h,b) in enumerate(lines):
    y=1.62+k*.7
    text(s,h,10.11,y,1.98,.48,20,'FFD67C',True,PP_ALIGN.RIGHT)
    text(s,b,1.15,y,8.7,.53,18,'FFF4D9',False,PP_ALIGN.RIGHT)
button(s,'ابدأ اللعبة',5.01,6.08,3.3,.65,'TITLE')
notes(s,'إعادة بناء مرئية وليست الملف الأصلي للبائع. شخصيات وعلامات Crash Bandicoot ملك أصحابها؛ نموذج للاستخدام الشخصي/التعليمي وليس ترخيصاً تجارياً. الخلفيات والعناصر مولّدة خصيصاً. المقدمة والبوابة مصيّرتان بثري جي إس. لا توجد نقاط تلقائية أو حماية امتحانات. يرجى مراجعة الأسئلة قبل استخدامها في صفك.')

for shape,target in pending:shape.click_action.target_slide=SL[target]
# Kiosk-style navigation: only authored links advance; Escape exits the show.
props_part=prs.part.part_related_by(RT.PRES_PROPS)
props=etree.fromstring(props_part.blob)
old=props.find('{'+P+'}showPr')
if old is not None:props.remove(old)
show=OxmlElement('p:showPr');show.set('loop','0');show.set('showAnimation','1');show.set('showNarration','1');show.set('useTimings','1');show.append(OxmlElement('p:kiosk'));show.append(OxmlElement('p:sldAll'));props.insert(0,show)
props_part._blob=etree.tostring(props,xml_declaration=True,encoding='UTF-8',standalone=True)
prs.save(ROOT/'Crash-Classroom.pptx')
(ROOT/'questions.json').write_text(json.dumps(Q,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created native PowerPoint:',len(prs.slides),'slides;',len(pending),'internal click targets')
