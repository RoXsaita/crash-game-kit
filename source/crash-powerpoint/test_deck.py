import unittest, zipfile, posixpath
from pathlib import Path
from lxml import etree
from pptx import Presentation
ROOT=Path(__file__).parent
FILE=ROOT/'Crash-Classroom.pptx'
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
class DeckTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue(FILE.exists(),'Native .pptx has not been built')
        self.prs=Presentation(FILE)
        self.slides={s.name:s for s in self.prs.slides}
    def test_native_deck_and_editable_content(self):
        self.assertAlmostEqual(self.prs.slide_width/self.prs.slide_height,16/9,places=3)
        for name in ['INTRO','TITLE','PORTAL','MENU','FINISH','TEACHER']:
            self.assertIn(name,self.slides)
        for i in range(8):
            s=self.slides[f'Q{i+1:02d}']
            self.assertTrue(any(sh.has_text_frame and sh.name=='QUESTION' for sh in s.shapes))
            self.assertTrue(s.notes_slide.notes_text_frame.text)
            self.assertIsNotNone(s._element.find('p:timing',NS))
    def test_menu_eight_native_click_targets(self):
        s=self.slides['MENU']
        crates=[sh for sh in s.shapes if sh.name.startswith('CRATE_')]
        self.assertEqual(len(crates),8)
        for i,sh in enumerate(crates):self.assertEqual(sh.click_action.target_slide.name,f'Q{i+1:02d}')
    def test_answer_navigation_and_feedback(self):
        for i in range(8):
            q=self.slides[f'Q{i+1:02d}']
            answers=[sh for sh in q.shapes if sh.name.startswith('ANSWER_')]
            self.assertIn(len(answers),[2,3])
            outcomes=[]
            for j,a in enumerate(answers):
                dest=a.click_action.target_slide
                self.assertEqual(dest.name,f'Q{i+1:02d}_A{j+1}')
                outcomes.extend([sh.name for sh in dest.shapes if sh.name in ['CORRECT','WRONG']])
                home=next(sh for sh in dest.shapes if sh.name=='HOME')
                self.assertEqual(home.click_action.target_slide.name,'MENU')
            self.assertEqual(outcomes.count('CORRECT'),1)
    def test_all_internal_targets_and_animation_ids(self):
        with zipfile.ZipFile(FILE) as z:
            files=set(z.namelist())
            for name in files:
                if name.endswith('.rels'):
                    root=etree.fromstring(z.read(name))
                    for rel in root:
                        self.assertNotEqual(rel.get('TargetMode'),'External')
                        base=posixpath.dirname(posixpath.dirname(name))
                        target=posixpath.normpath(posixpath.join(base,rel.get('Target'))).lstrip('/')
                        self.assertIn(target,files,(name,target))
            self.assertFalse(any('vbaProject' in name for name in files))
            for s in self.prs.slides:
                ids=[e.get('id') for e in s._element.findall('.//p:cTn',NS)]
                self.assertEqual(len(ids),len(set(ids)))
                shapes={str(sh.shape_id) for sh in s.shapes}
                for target in s._element.findall('.//p:spTgt',NS):self.assertIn(target.get('spid'),shapes)
    def test_embedded_animations_audio_and_navigation_safety(self):
        with zipfile.ZipFile(FILE) as z:
            self.assertGreaterEqual(sum(n.endswith('.mp4') for n in z.namelist()),2)
            self.assertGreaterEqual(sum(n.endswith('.wav') for n in z.namelist()),2)
        for name in ['INTRO','PORTAL']:
            s=self.slides[name]
            self.assertIsNotNone(s._element.find('p:transition',NS).get('advTm'))
            conditions=s._element.findall('.//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond',NS)
            self.assertTrue(conditions)
            self.assertTrue(all(c.get('delay')=='0' for c in conditions))
        for name,s in self.slides.items():
            self.assertEqual(s._element.find('p:transition',NS).get('advClick'),'0')
if __name__=='__main__':unittest.main(verbosity=2)
